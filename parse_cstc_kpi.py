"""
Parses CSTC sheets in Matrix for Ranking System of PSTDs.xlsx and outputs
cstc-kpi-scores.csv using the same 54-subrow scheme as parse_excel_kpi.py.

Sheet names like 'CSTC-Davao City' map to province name 'Davao City'
(strip 'CSTC-' prefix). Skips 'CSTC-Rename*' placeholders.
"""
import openpyxl
import csv
import re

from parse_excel_kpi import (
    SUBROW_KPI, YEAR_COLS, find_subrow_rows,
    normalize_value, sum_values,
)

EXCEL_PATH  = r'c:\Herd\DOST-Provincial-Director-Ranking-System\Matrix for Ranking System of PSTDs.xlsx'
OUTPUT_PATH = r'c:\Herd\DOST-Provincial-Director-Ranking-System\storage\app\private\data\cstc-kpi-scores.csv'


def parse_sheet(ws, province_name):
    rows_data = list(ws.iter_rows(values_only=True))
    row_map   = find_subrow_rows(rows_data)
    records   = []

    for subrow_id, row_idx in row_map.items():
        kpi_id = SUBROW_KPI.get(subrow_id)
        if kpi_id is None:
            continue

        indices = row_idx if isinstance(row_idx, list) else [row_idx]
        rows    = [rows_data[i] for i in indices if i < len(rows_data)]
        if not rows:
            continue

        for year, (tcol, acol) in YEAR_COLS.items():
            target_cells = [r[tcol] if tcol < len(r) else None for r in rows]
            acc_cells    = [r[acol] if acol < len(r) else None for r in rows]
            if len(indices) > 1:
                target       = sum_values(target_cells)
                accomplished = sum_values(acc_cells)
            else:
                target       = normalize_value(target_cells[0])
                accomplished = normalize_value(acc_cells[0])
            if target or accomplished:
                records.append({
                    'province':     province_name,
                    'kpi_id':       kpi_id,
                    'subrow_id':    subrow_id,
                    'year':         year,
                    'target':       target,
                    'accomplished': accomplished,
                })
    return records


def main():
    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    all_records = []

    for sheet_name in wb.sheetnames:
        stripped = sheet_name.strip()
        if not stripped.startswith('CSTC-'):
            continue
        # Skip placeholder sheets like CSTC-Rename6, CSTC-Rename7, ...
        province_name = stripped[len('CSTC-'):].strip()
        if re.match(r'(?i)^rename\d*$', province_name):
            print(f'  skip placeholder: {sheet_name}')
            continue

        ws      = wb[sheet_name]
        records = parse_sheet(ws, province_name)
        if records:
            all_records.extend(records)
            print(f'  {province_name}: {len(records)} records')
        else:
            print(f'  {province_name}: no data')

    print(f'\nTotal CSTC records: {len(all_records)}')

    with open(OUTPUT_PATH, 'w', newline='', encoding='utf-8') as f:
        writer = csv.DictWriter(f, fieldnames=['province', 'kpi_id', 'subrow_id', 'year', 'target', 'accomplished'])
        writer.writeheader()
        writer.writerows(all_records)

    print(f'Written to {OUTPUT_PATH}')


if __name__ == '__main__':
    main()
