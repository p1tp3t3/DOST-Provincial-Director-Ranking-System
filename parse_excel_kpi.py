"""
Parses Matrix for Ranking System of PSTDs.xlsx and outputs kpi-scores.csv.
54-indicator scheme matching the sheet's numbered items 1-54. Sub-categories
(a, b, c under a parent indicator) are SUMMED into the parent row:
    #17 Firms graduated = Micro->Small + Small->Medium + Medium->Large
    #25 DRRM measures   = Activities + IEC Materials
    #53 Linkages        = MOA + MOU
    #54 External funds  = DOST Councils + Other NGAs + Funding Partners

Output CSV columns: province, subrow_id, kpi_id, year, target, accomplished
Empty string = NULL (no data).
"""
import openpyxl
import csv

EXCEL_PATH  = r'c:\Herd\DOST-Provincial-Director-Ranking-System\Matrix for Ranking System of PSTDs.xlsx'
OUTPUT_PATH = r'c:\Herd\DOST-Provincial-Director-Ranking-System\storage\app\private\data\kpi-scores.csv'

# KPI (outcome) ID each subrow belongs to — 54-indicator scheme
SUBROW_KPI = {
    **{i: 1 for i in [1, 2]},
    **{i: 2 for i in [3, 4, 5, 6, 7, 8]},
    **{i: 3 for i in [9]},
    **{i: 4 for i in range(10, 25)},   # 10-24 (15 indicators)
    **{i: 5 for i in [25, 26]},
    **{i: 6 for i in [27, 28]},
    **{i: 7 for i in range(29, 55)},   # 29-54 (26 indicators)
}

# Year → (target_col, accomplished_col) 0-indexed
YEAR_COLS = {2022: (2, 3), 2023: (4, 5), 2024: (6, 7), 2025: (8, 9)}

SKIP_SHEETS = {'Province Directory (Quick Acces', 'Sheet6', 'Sheet8'}


def normalize_value(v):
    if v is None:
        return ''
    s = str(v).strip()
    if s.lower() in ('-', '', 'no target', 'none', 'na', 'n/a', 'x'):
        return ''
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    if isinstance(v, float):
        return str(round(v, 6))
    if isinstance(v, int):
        return str(v)
    return s


def to_number(v):
    """Convert a cell value to a float, or None if blank/non-numeric."""
    if v is None:
        return None
    s = str(v).strip()
    if s.lower() in ('-', '', 'no target', 'none', 'na', 'n/a', 'x'):
        return None
    try:
        return float(s)
    except (ValueError, TypeError):
        return None


def sum_values(values):
    """Sum a list of cell values. Returns '' if ALL are blank, else the sum."""
    nums = [to_number(v) for v in values]
    nums = [n for n in nums if n is not None]
    if not nums:
        return ''
    total = sum(nums)
    if total == int(total):
        return str(int(total))
    return str(round(total, 6))


def find_subrow_rows(rows_data):
    """
    Scans the sheet and returns {subrow_id: row_index_or_list}.
    Single-row indicators → int. Summed parent indicators → list of int.
    Uses label-text matching so it works regardless of row offsets per sheet.
    """
    n = len(rows_data)

    def lbl(i):
        return str(rows_data[i][1] or '').strip().lower() if i < n else ''

    def find_all(fragment):
        return sorted(i for i in range(n) if fragment in lbl(i))

    def first(fragment):
        r = find_all(fragment)
        return r[0] if r else None

    m = {}

    # ── Outcome 1: Innovation Stimulated ────────────────────────────────────
    r = first('r&d proposals evaluated') or first('r&d proposals')
    if r is not None: m[1] = r

    r = first('collaborative r&d')
    if r is not None: m[2] = r

    # ── Outcome 2: Technology Adoption ──────────────────────────────────────
    r = first('knowledge and technologies transferred')
    if r is not None: m[3] = r

    for i in find_all('technologies promoted'):
        if 'adopted' not in lbl(i): m[4] = i; break

    for i in find_all('technologies adopted'):
        m[5] = i; break

    r = first('technology adoptors') or first('technology adopters')
    if r is not None: m[6] = r

    r = first('promotion activities for dost')
    if r is not None: m[7] = r

    r = first('technology transfer and commercialization')
    if r is not None: m[8] = r

    # ── Outcome 3: Foster STI Culture ───────────────────────────────────────
    r = first('s&t promotional activities')
    if r is not None: m[9] = r

    # ── Outcome 4: Productivity and Efficiency ───────────────────────────────
    r = first('s&t interventions provided')
    if r is not None: m[10] = r

    for i in find_all('start-ups'):
        if 'spin-offs' in lbl(i) or 'enterprises' in lbl(i): m[11] = i; break

    for i in find_all('customers assisted'):
        m[12] = i; break

    for i in find_all('msmes'):
        m[13] = i; break

    r = first('employments generated')
    if r is not None: m[14] = r

    r = first('gross sales generated')
    if r is not None: m[15] = r

    r = first('increase in productivity')
    if r is not None: m[16] = r

    # #17: SUM of Micro→Small + Small→Medium + Medium→Large
    firms_grad = []
    r = first('micro to small')
    if r is not None: firms_grad.append(r)
    r = first('small to medium')
    if r is not None: firms_grad.append(r)
    r = first('medium to large')
    if r is not None: firms_grad.append(r)
    if firms_grad: m[17] = sorted(set(firms_grad))

    r = first('firms that became exporters')
    if r is not None: m[18] = r

    r = first('innovation hubs')
    if r is not None: m[19] = r

    for i in find_all('communities assisted'):
        if 'technologies' not in lbl(i): m[20] = i; break

    r = first('technologies deployed to communities')
    if r is not None: m[21] = r

    r = first('number of beneficiaries')
    if r is not None: m[22] = r

    r = first('samples referred by psto')
    if r is not None: m[23] = r

    r = first('testing and/or calibration services')
    if r is not None: m[24] = r

    # ── Outcome 5: Resiliency ────────────────────────────────────────────────
    # #25: SUM of Activities + IEC Materials under DRRM measures parent
    drrm_anchor = first('measures on disaster risk reduction')
    drrm_rows = []
    if drrm_anchor is not None:
        for j in range(drrm_anchor + 1, min(drrm_anchor + 6, n)):
            l = lbl(j)
            if 'iec material' in l:
                drrm_rows.append(j)
            elif 'activities' in l and j != drrm_anchor:
                raw_l = l.strip()
                if len(raw_l) < 50 or raw_l.startswith('a.'):
                    drrm_rows.append(j)
    if drrm_rows: m[25] = sorted(set(drrm_rows))

    r = first('drrm-related collaboration')
    if r is not None: m[26] = r

    # ── Outcome 6: Effective STI Governance ─────────────────────────────────
    r = first('programs/projects with risk analysis') or first('% of programs')
    if r is not None: m[27] = r

    r = first('work-process innovation')
    if r is not None: m[28] = r

    # ── Outcome 7: Other Indicators ─────────────────────────────────────────
    r = first('setup projects endorsed')
    if r is not None: m[29] = r

    for i in find_all('setup projects approved'):
        if 'value' not in lbl(i): m[30] = i; break

    r = first('value of setup projects approved')
    if r is not None: m[31] = r

    r = first('setup projects implemented')
    if r is not None: m[32] = r

    r = first('ongoing setup projects')
    if r is not None: m[33] = r

    r = first('completed setup projects')
    if r is not None: m[34] = r

    r = first('delinquent setup projects')
    if r is not None: m[35] = r

    # SETUP Refund Rate appears twice — first occurrence → 36, second → 51
    rr_rows = find_all('setup refund rate')
    if len(rr_rows) > 0: m[36] = rr_rows[0]
    if len(rr_rows) > 1: m[51] = rr_rows[1]

    # SETUP Refunded Amount appears twice — first → 37, second → 52
    ra_rows = find_all('setup refunded amount')
    if len(ra_rows) > 0: m[37] = ra_rows[0]
    if len(ra_rows) > 1: m[52] = ra_rows[1]

    r = first('customer satisfaction rating') or first('customer satisfaction')
    if r is not None: m[38] = r

    for i in find_all('gia proposal approved'):
        if 'value' not in lbl(i): m[39] = i; break

    r = first('value of gia projects approved') or first('value of gia project')
    if r is not None: m[40] = r

    r = first('gia projects implemented')
    if r is not None: m[41] = r

    r = first('ongoing gia projects')
    if r is not None: m[42] = r

    r = first('completed gia projects')
    if r is not None: m[43] = r

    for i in find_all('cest proposal approved'):
        if 'value' not in lbl(i): m[44] = i; break

    r = first('value of cest proposal')
    if r is not None: m[45] = r

    r = first('cest projects implemented')
    if r is not None: m[46] = r

    r = first('ongoing cest projects')
    if r is not None: m[47] = r

    r = first('completed cest projects')
    if r is not None: m[48] = r

    r = first('draft press releases')
    if r is not None: m[49] = r

    r = first('facebook post')
    if r is not None: m[50] = r

    # #53: SUM of MOA + MOU under "Number of linkages established"
    linkages_anchor = first('number of linkages established')
    linkage_rows = []
    if linkages_anchor is not None:
        for j in range(linkages_anchor + 1, min(linkages_anchor + 5, n)):
            l = lbl(j)
            if 'moa' in l or 'mou' in l:
                linkage_rows.append(j)
    if linkage_rows: m[53] = sorted(set(linkage_rows))

    # #54: SUM of DOST Councils + Other NGAs + Funding Partners
    funds_anchor = first('amount of external funds')
    fund_rows = []
    if funds_anchor is not None:
        for j in range(funds_anchor + 1, min(funds_anchor + 8, n)):
            l = lbl(j)
            if 'dost councils' in l or 'other ngas' in l or 'funding partner' in l:
                fund_rows.append(j)
    if fund_rows: m[54] = sorted(set(fund_rows))

    return m


def parse_sheet(ws, province_name):
    rows_data = list(ws.iter_rows(values_only=True))
    row_map   = find_subrow_rows(rows_data)
    records   = []

    for subrow_id, row_idx in row_map.items():
        kpi_id = SUBROW_KPI.get(subrow_id)
        if kpi_id is None:
            continue

        indices = row_idx if isinstance(row_idx, list) else [row_idx]
        rows    = [rows_data[i] for i in indices if i < len(rows_data)]
        if not rows:
            continue

        for year, (tcol, acol) in YEAR_COLS.items():
            target_cells = [r[tcol] if tcol < len(r) else None for r in rows]
            acc_cells    = [r[acol] if acol < len(r) else None for r in rows]
            if len(indices) > 1:
                target       = sum_values(target_cells)
                accomplished = sum_values(acc_cells)
            else:
                target       = normalize_value(target_cells[0])
                accomplished = normalize_value(acc_cells[0])
            if target or accomplished:
                records.append({
                    'province':     province_name,
                    'subrow_id':    subrow_id,
                    'kpi_id':       kpi_id,
                    'year':         year,
                    'target':       target,
                    'accomplished': accomplished,
                })
    return records


def main():
    wb          = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    all_records = []
    skipped     = []

    for sheet_name in wb.sheetnames:
        if sheet_name in SKIP_SHEETS or sheet_name.startswith('CSTC'):
            continue
        ws      = wb[sheet_name]
        records = parse_sheet(ws, sheet_name)
        if records:
            all_records.extend(records)
            print(f'  {sheet_name}: {len(records)} records')
        else:
            skipped.append(sheet_name)
            print(f'  {sheet_name}: no data (skipped)')

    print(f'\nTotal records: {len(all_records)}')
    if skipped:
        print(f'Sheets with no data: {skipped}')

    with open(OUTPUT_PATH, 'w', newline='', encoding='utf-8') as f:
        writer = csv.DictWriter(f, fieldnames=['province', 'subrow_id', 'kpi_id', 'year', 'target', 'accomplished'])
        writer.writeheader()
        writer.writerows(all_records)

    print(f'Written to {OUTPUT_PATH}')


if __name__ == '__main__':
    main()
